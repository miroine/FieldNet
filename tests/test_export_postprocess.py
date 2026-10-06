import io, json, subprocess, sys, time, zipfile, urllib.request, socket, pytest, pandas as pd
from tests.hubfix import demo_hub
from network import exporters as ex, postprocess as pp


def _ds(): return demo_hub()[4]


def test_csv_zip_roundtrip():
    h = _ds(); z = zipfile.ZipFile(io.BytesIO(ex.to_csv_zip(h.datasets, h.meta, h.info)))
    assert 'manifest.json' in z.namelist() and 'annual_field.csv' in z.namelist()
    d = pd.read_csv(io.BytesIO(z.read('annual_field.csv')))
    assert d['Oil [Sm3]'].sum() == pytest.approx(h.datasets['annual_field']['Oil [Sm3]'].sum())


def test_excel_sheets_and_readme():
    import openpyxl
    h = _ds(); wb = openpyxl.load_workbook(io.BytesIO(ex.to_excel(h.datasets, h.meta, h.info)))
    assert wb.sheetnames[0] == 'README' and len(wb.sheetnames) == len(h.datasets) + 1 and all(len(s) <= 31 for s in wb.sheetnames)


def test_excel_unique_long_names():
    d = {'a' * 40: pd.DataFrame({'x': [1]}), 'a' * 40 + 'b': pd.DataFrame({'x': [2]})}
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(ex.to_excel(d))); assert len(set(wb.sheetnames)) == 3


def test_json_export():
    h = _ds(); j = json.loads(ex.to_json(h.datasets, h.meta, h.info)); assert 'annual_field' in j['data'] and j['manifest']['generator'] == 'FieldNet'


def test_stea_table_units_and_totals():
    h = _ds(); t = ex.stea_table(h.datasets)
    oil = t[t['Series'] == 'Oil production'].iloc[0]
    assert oil['Unit'] == 'MSm3' and oil['Total'] == pytest.approx(h.datasets['annual_field']['Oil [Sm3]'].sum() / 1e6)
    assert not t.attrs['skipped']


def test_stea_custom_mapping_and_missing():
    h = _ds(); m = ex.mapping_from_table(pd.DataFrame([{'Series': 'X', 'Dataset': 'annual_field', 'Column': 'Oil [Sm3]', 'Scale': 0.001, 'Unit': 'k'}, {'Series': 'Y', 'Dataset': 'nope', 'Column': 'c'}]))
    t = ex.stea_table(h.datasets, m); assert list(t['Series']) == ['X'] and t.attrs['skipped']
    assert ';' in ex.stea_csv(h.datasets)


def test_api_bundle_serves(tmp_path):
    h = _ds(); z = zipfile.ZipFile(io.BytesIO(ex.api_bundle(h.datasets, h.meta, h.info))); z.extractall(tmp_path)
    s = socket.socket(); s.bind(('127.0.0.1', 0)); port = s.getsockname()[1]; s.close()
    pr = subprocess.Popen([sys.executable, str(tmp_path / 'serve.py'), str(port)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        for _ in range(50):
            try: lst = json.load(urllib.request.urlopen(f'http://127.0.0.1:{port}/datasets')); break
            except Exception: time.sleep(0.1)
        assert 'annual_field' in lst
        rows = json.load(urllib.request.urlopen(f'http://127.0.0.1:{port}/datasets/annual_field'))
        assert len(rows) == len(h.datasets['annual_field'])
        assert b'Year' in urllib.request.urlopen(f'http://127.0.0.1:{port}/datasets/annual_field?format=csv').read()
        with pytest.raises(Exception): urllib.request.urlopen(f'http://127.0.0.1:{port}/datasets/../manifest.json_x')
    finally: pr.kill()


def test_postprocess_runs_and_does_not_mutate():
    h = _ds(); before = h.datasets['annual_field'].copy()
    t, log = pp.run("d = ds['annual_field'].copy()\nd['Oil [MSm3]'] = d['Oil [Sm3]']/1e6\nout['x'] = d\nprint('ok')", h.datasets)
    assert 'Oil [MSm3]' in t['x'] and log.strip() == 'ok'
    pp.run("ds['annual_field']['Oil [Sm3]'] = 0\nout['y'] = ds['annual_field']", h.datasets)
    assert h.datasets['annual_field'].equals(before)


@pytest.mark.parametrize('src', ['import os', 'open("x")', "__import__('os')", "ds.__class__", "pd.read_csv('x')", "eval('1')", "getattr(pd,'x')", "class A: pass", "x = (1).__class__"])
def test_postprocess_rejects(src):
    with pytest.raises(pp.ScriptError): pp.run(src + "\nout['a']=pd.DataFrame({'a':[1]})", {})


def test_postprocess_errors_readable():
    with pytest.raises(pp.ScriptError, match='KeyError'): pp.run("out['a']=ds['missing']", {})
    with pytest.raises(pp.ScriptError, match='no table'): pp.run("x=1", {})
    with pytest.raises(pp.ScriptError, match='syntax'): pp.run("def (", {})


def test_postprocess_examples_work():
    h = _ds()
    for name, src in pp.EXAMPLES.items():
        t, _ = pp.run(src, h.datasets, {}); assert t, name
